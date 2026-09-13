import os
import re
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, scrolledtext
from datetime import datetime
import docx
from PyPDF2 import PdfReader
import traceback
import platform


def resource_path(relative_path):
    """获取打包后资源的绝对路径"""
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


class FileRenameApp:
    def __init__(self, root):
        self.root = root
        self.root.title("多功能文件重命名工具")
        self.root.geometry("900x820")
        self.font = ('Microsoft YaHei UI', 10)

        # 存储预览结果
        self.preview_results = []
        self.is_single_file = False
        self.current_mode = "batch"  # 默认选择批量重命名
        self.drag_data = {"index": None, "item": None}  # 用于拖动排序的数据
        self.file_list = []  # 存储当前文件列表
        self.open_files_check_var = tk.BooleanVar(value=True)  # 新增：是否检查打开的文件

        # 创建界面元素
        self.create_widgets()
        # 绑定 F5 快捷键刷新
        self.root.bind('<F5>', lambda e: self.refresh_folder() if self.current_mode == "batch" else None)
        print("[INFO] 界面初始化完成")

    def create_widgets(self):
        # 模式选择框架
        mode_frame = tk.Frame(self.root, padx=10, pady=5)
        mode_frame.pack(fill=tk.X)

        self.mode_var = tk.StringVar(value="batch")
        batch_radio = tk.Radiobutton(mode_frame, text="批量重命名",
                                     variable=self.mode_var, value="batch",
                                     font=self.font, command=self.change_mode)
        batch_radio.pack(side=tk.LEFT, padx=10)

        single_radio = tk.Radiobutton(mode_frame, text="单个文件重命名",
                                      variable=self.mode_var, value="single",
                                      font=self.font, command=self.change_mode)
        single_radio.pack(side=tk.LEFT, padx=10)

        title_radio = tk.Radiobutton(mode_frame, text="按文件标题重命名",
                                     variable=self.mode_var, value="title",
                                     font=self.font, command=self.change_mode)
        title_radio.pack(side=tk.LEFT, padx=10)

        # 路径选择框架
        path_frame = tk.Frame(self.root, padx=10, pady=10)
        path_frame.pack(fill=tk.X)

        tk.Label(path_frame, text="文件/文件夹路径:", font=self.font).pack(side=tk.LEFT)

        self.path_entry = tk.Entry(path_frame, width=50, font=self.font)
        self.path_entry.pack(side=tk.LEFT, padx=5)

        self.folder_btn = tk.Button(path_frame, text="浏览文件夹", font=self.font,
                                    command=self.browse_folder)
        self.folder_btn.pack(side=tk.LEFT, padx=5)

        self.file_btn = tk.Button(path_frame, text="浏览文件", font=self.font,
                                  command=self.browse_file)
        self.file_btn.pack(side=tk.LEFT, padx=5)

        # 👇 新增：刷新文件夹按钮
        self.refresh_btn = tk.Button(path_frame, text="刷新文件夹", font=self.font,
                                     command=self.refresh_folder, state=tk.DISABLED)
        self.refresh_btn.pack(side=tk.LEFT, padx=5)

        self.mode_label = tk.Label(path_frame, text="批量处理模式", font=self.font, fg="blue")
        self.mode_label.pack(side=tk.LEFT, padx=10)

        # 预设命名格式框架
        self.preset_frame = tk.Frame(self.root, padx=10, pady=5)
        self.preset_frame.pack(fill=tk.X)

        tk.Label(self.preset_frame, text="预设命名格式:", font=self.font).pack(side=tk.LEFT, padx=5)

        self.preset_var = tk.StringVar(value="不使用预设")
        self.preset_combobox = ttk.Combobox(
            self.preset_frame,
            textvariable=self.preset_var,
            values=[
                "不使用预设",
                "序号模式 (001, 002...)",
                "日期模式 (YYYYMMDD)",
                "时间戳模式 (YYYYMMDD_HHMMSS)",
                "序号+原文件名",
                "日期+原文件名",
                "原文件名+序号",
                "序号_原文件名_日期"
            ],
            width=25,
            font=self.font
        )
        self.preset_combobox.pack(side=tk.LEFT, padx=5)
        self.preset_combobox.bind("<<ComboboxSelected>>", self.on_preset_change)

        # 序号设置框架
        self.numbering_frame = tk.Frame(self.root, padx=10, pady=0)

        self.start_num_var = tk.IntVar(value=1)
        tk.Label(self.numbering_frame, text="起始序号:", font=self.font).pack(side=tk.LEFT, padx=5)
        start_num_entry = tk.Entry(self.numbering_frame, textvariable=self.start_num_var, width=5, font=self.font)
        start_num_entry.pack(side=tk.LEFT, padx=5)
        start_num_entry.bind("<KeyRelease>", self.on_numbering_change)

        self.digits_var = tk.IntVar(value=2)
        tk.Label(self.numbering_frame, text="数字位数:", font=self.font).pack(side=tk.LEFT, padx=5)
        digits_entry = tk.Entry(self.numbering_frame, textvariable=self.digits_var, width=5, font=self.font)
        digits_entry.pack(side=tk.LEFT, padx=5)
        digits_entry.bind("<KeyRelease>", self.on_numbering_change)

        # 日期格式设置框架
        self.date_format_frame = tk.Frame(self.root, padx=10, pady=0)

        self.date_format_var = tk.StringVar(value="YYYYMMDD")
        date_format_label = tk.Label(self.date_format_frame, text="日期格式:", font=self.font)
        date_format_label.pack(side=tk.LEFT, padx=5)

        date_format_combobox = ttk.Combobox(
            self.date_format_frame,
            textvariable=self.date_format_var,
            values=["YYYYMMDD", "YYYY-MM-DD", "YYMMDD", "YY-MM-DD"],
            width=10,
            font=self.font
        )
        date_format_combobox.pack(side=tk.LEFT, padx=5)
        date_format_combobox.bind("<<ComboboxSelected>>", self.on_date_format_change)

        # 智能识别已有格式文件的选项
        self.ignore_existing_format_var = tk.BooleanVar(value=False)
        ignore_check = tk.Checkbutton(
            self.date_format_frame,
            text="忽略已符合格式的文件",
            variable=self.ignore_existing_format_var,
            font=self.font,
            command=self.on_ignore_existing_change
        )
        ignore_check.pack(side=tk.LEFT, padx=10)

        # 检查打开文件的选项
        open_files_check = tk.Checkbutton(
            self.date_format_frame,
            text="检查已打开的文件",
            variable=self.open_files_check_var,
            font=self.font
        )
        open_files_check.pack(side=tk.LEFT, padx=10)

        # 标题重命名选项框架
        self.title_options_frame = tk.Frame(self.root, padx=10, pady=5)
        self.title_options_frame.pack(fill=tk.X)

        self.auto_rename_title_var = tk.BooleanVar(value=False)
        auto_rename_title_check = tk.Checkbutton(self.title_options_frame, text="自动重命名",
                                                 variable=self.auto_rename_title_var, font=self.font)
        auto_rename_title_check.pack(side=tk.LEFT, padx=5)

        self.title_prefix_var = tk.StringVar(value="")
        prefix_label = tk.Label(self.title_options_frame, text="前缀:", font=self.font)
        prefix_label.pack(side=tk.LEFT, padx=5)
        prefix_entry = tk.Entry(self.title_options_frame, textvariable=self.title_prefix_var, width=10, font=self.font)
        prefix_entry.pack(side=tk.LEFT, padx=5)

        self.title_suffix_var = tk.StringVar(value="")
        suffix_label = tk.Label(self.title_options_frame, text="后缀:", font=self.font)
        suffix_label.pack(side=tk.LEFT, padx=5)
        suffix_entry = tk.Entry(self.title_options_frame, textvariable=self.title_suffix_var, width=10, font=self.font)
        suffix_entry.pack(side=tk.LEFT, padx=5)

        # 文件格式后缀修改框架
        self.extension_frame = tk.Frame(self.root, padx=10, pady=5)
        self.extension_frame.pack(fill=tk.X)

        tk.Label(self.extension_frame, text="修改文件后缀:", font=self.font).pack(side=tk.LEFT, padx=5)

        self.common_extensions = [
            "不修改", "txt", "pdf", "docx", "doc", "xlsx", "xls", "pptx", "ppt",
            "jpg", "jpeg", "png", "gif", "bmp", "svg", "mp4", "avi", "mov",
            "mp3", "wav", "zip", "rar", "7z", "html", "htm", "json", "csv"
        ]

        self.extension_var = tk.StringVar(value="不修改")
        self.extension_combobox = ttk.Combobox(
            self.extension_frame,
            textvariable=self.extension_var,
            values=self.common_extensions,
            width=10,
            font=self.font
        )
        self.extension_combobox.pack(side=tk.LEFT, padx=5)
        self.extension_combobox.bind("<<ComboboxSelected>>", self.on_extension_change)

        # 自定义后缀输入框
        tk.Label(self.extension_frame, text="自定义:", font=self.font).pack(side=tk.LEFT, padx=5)
        self.custom_ext_var = tk.StringVar(value="")
        self.custom_ext_entry = tk.Entry(
            self.extension_frame,
            textvariable=self.custom_ext_var,
            width=10,
            font=self.font
        )
        self.custom_ext_entry.pack(side=tk.LEFT, padx=5)
        self.custom_ext_entry.bind("<KeyRelease>", self.on_custom_ext_change)

        # 预览结果框架
        preview_frame = tk.Frame(self.root, padx=10, pady=10)
        preview_frame.pack(fill=tk.BOTH, expand=True)

        # 左侧 - 原始文件
        left_frame = tk.Frame(preview_frame)
        left_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 5))

        self.original_label = tk.Label(left_frame, text="原始文件列表:", font=self.font, fg="blue")
        self.original_label.pack(anchor=tk.W)

        # 使用 Listbox 替代 ScrolledText，支持拖动排序
        self.original_files_listbox = tk.Listbox(
            left_frame,
            width=40,
            height=5,
            font=self.font,
            selectmode=tk.EXTENDED,
            activestyle=tk.NONE
        )
        self.original_files_listbox.pack(fill=tk.BOTH, expand=True, pady=5)

        # 绑定拖动排序事件
        self.original_files_listbox.bind("<Button-1>", self.on_drag_start)
        self.original_files_listbox.bind("<B1-Motion>", self.on_drag_motion)
        self.original_files_listbox.bind("<ButtonRelease-1>", self.on_drag_end)

        # 右侧 - 提取结果
        right_frame = tk.Frame(preview_frame)
        right_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(5, 0))

        self.preview_label = tk.Label(right_frame, text="重命名预览:", font=self.font, fg="green")
        self.preview_label.pack(anchor=tk.W)

        # 使用 Listbox 替代 ScrolledText，支持拖动排序
        self.preview_listbox = tk.Listbox(
            right_frame,
            width=40,
            height=5,
            font=self.font,
            selectmode=tk.EXTENDED,
            activestyle=tk.NONE
        )
        self.preview_listbox.pack(fill=tk.BOTH, expand=True, pady=5)

        # 绑定拖动排序事件
        self.preview_listbox.bind("<Button-1>", self.on_drag_start)
        self.preview_listbox.bind("<B1-Motion>", self.on_drag_motion)
        self.preview_listbox.bind("<ButtonRelease-1>", self.on_drag_end)

        # 按钮框架
        btn_frame = tk.Frame(self.root, padx=10, pady=10)
        btn_frame.pack(fill=tk.X)

        self.extract_title_btn = tk.Button(btn_frame, text="提取标题", font=self.font,
                                           command=self.extract_title)
        self.extract_title_btn.pack_forget()

        self.execute_btn = tk.Button(btn_frame, text="执行重命名", font=self.font,
                                     command=self.execute_rename, state=tk.DISABLED)
        self.execute_btn.pack(side=tk.LEFT, padx=5)

        self.remove_format_btn = tk.Button(btn_frame, text="取消格式", font=self.font,
                                           command=self.remove_format, state=tk.DISABLED)
        self.remove_format_btn.pack(side=tk.LEFT, padx=5)

        # 状态框架
        status_frame = tk.Frame(self.root, padx=10, pady=5)
        status_frame.pack(fill=tk.X)

        self.status_var = tk.StringVar(value="就绪")
        status_label = tk.Label(status_frame, textvariable=self.status_var, font=self.font, fg="blue")
        status_label.pack(anchor=tk.W)

        # 初始化界面
        self.change_mode()
        print("[INFO] 界面组件加载完成")

    def refresh_folder(self):
        """刷新当前文件夹的文件列表与预览"""
        folder_path = self.path_entry.get().strip()
        if not folder_path:
            messagebox.showwarning("提示", "请先选择或输入文件夹路径")
            return
        if not os.path.isdir(folder_path):
            messagebox.showerror("错误", "路径无效或不是文件夹，请重新选择")
            return

        self.status_var.set("正在刷新文件夹...")
        self.root.update_idletasks()
        self.show_batch_files(folder_path)
        self.status_var.set(f"已刷新文件夹: {os.path.basename(folder_path)}")
        print(f"[INFO] 文件夹刷新完成: {folder_path}")

    def on_drag_start(self, event):
        """开始拖动"""
        widget = event.widget
        self.drag_data["index"] = widget.nearest(event.y)
        self.drag_data["item"] = widget.get(self.drag_data["index"])

        # 高亮显示当前选中的项
        widget.selection_clear(0, tk.END)
        widget.selection_set(self.drag_data["index"])
        widget.see(self.drag_data["index"])

    def on_drag_motion(self, event):
        """拖动过程中"""
        widget = event.widget
        index = widget.nearest(event.y)

        if index != self.drag_data["index"]:
            # 删除原位置的项
            widget.delete(self.drag_data["index"])
            # 在新位置插入项
            widget.insert(index, self.drag_data["item"])
            # 更新选中的索引
            self.drag_data["index"] = index
            # 更新高亮显示
            widget.selection_clear(0, tk.END)
            widget.selection_set(index)
            widget.see(index)

    def on_drag_end(self, event):
        """拖动结束"""
        widget = event.widget
        new_index = widget.nearest(event.y)

        # 如果是原始文件列表框，更新预览列表
        if widget == self.original_files_listbox and self.preview_results:
            # 获取新的文件顺序
            new_files = [widget.get(i) for i in range(widget.size())]

            # 更新文件列表
            self.file_list = new_files

            folder_path = self.path_entry.get()

            # 重新生成预览结果，确保序号正确
            self.generate_batch_preview(new_files, folder_path)

    def change_mode(self):
        self.current_mode = self.mode_var.get()
        print(f"[INFO] 切换到{self.current_mode}模式")

        # 👇 根据模式控制刷新按钮状态
        if self.current_mode == "batch":
            self.refresh_btn.config(state=tk.NORMAL)
        else:
            self.refresh_btn.config(state=tk.DISABLED)

        if self.current_mode == "batch":
            self.is_single_file = False
            self.mode_label.config(text="批量处理模式")
            self.folder_btn.config(text="浏览文件夹", command=self.browse_folder)
            self.extract_title_btn.pack_forget()
            self.original_label.config(text="原始文件列表:")
            self.preview_label.config(text="重命名预览:")

        elif self.current_mode == "single":
            self.is_single_file = True
            self.mode_label.config(text="单个文件模式")
            self.folder_btn.config(text="浏览文件", command=self.browse_file)
            self.extract_title_btn.pack_forget()
            self.original_label.config(text="原始文件:")
            self.preview_label.config(text="重命名预览:")

        elif self.current_mode == "title":
            self.is_single_file = True
            self.mode_label.config(text="按标题重命名模式")
            self.folder_btn.config(text="浏览文件", command=self.browse_file)
            self.extract_title_btn.pack(side=tk.LEFT, padx=5)
            self.original_label.config(text="原始文件名:")
            self.preview_label.config(text="提取结果:")

        self.preset_frame.pack(fill=tk.X)
        self.update_numbering_frame_visibility()
        self.update_date_format_frame_visibility()
        self.title_options_frame.pack(fill=tk.X)
        self.extension_frame.pack(fill=tk.X)

        self.path_entry.delete(0, tk.END)
        self.original_files_listbox.delete(0, tk.END)
        self.preview_listbox.delete(0, tk.END)
        self.execute_btn.config(state=tk.DISABLED)
        self.remove_format_btn.config(state=tk.DISABLED)
        self.status_var.set("就绪")
        print("[INFO] 模式切换完成")

    def update_numbering_frame_visibility(self):
        """根据预设选择显示或隐藏序号设置框架"""
        preset = self.preset_var.get()
        if preset in ["序号模式 (001, 002...)", "序号+原文件名", "原文件名+序号", "序号_原文件名_日期"]:
            self.numbering_frame.pack(fill=tk.X, pady=2)
        else:
            self.numbering_frame.pack_forget()

    def update_date_format_frame_visibility(self):
        """根据预设选择显示或隐藏日期格式设置框架"""
        preset = self.preset_var.get()
        if preset in ["日期模式 (YYYYMMDD)", "日期+原文件名", "序号_原文件名_日期"]:
            self.date_format_frame.pack(fill=tk.X, pady=2)
        else:
            self.date_format_frame.pack_forget()

    def on_preset_change(self, event=None):
        """预设命名格式选择变化时的处理"""
        self.update_numbering_frame_visibility()
        self.update_date_format_frame_visibility()

        # 如果是批量模式，自动更新预览
        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                self.generate_batch_preview(self.file_list, folder_path)
        elif self.current_mode == "single" and self.path_entry.get():
            self.generate_single_preview(self.path_entry.get())

    def on_numbering_change(self, event=None):
        """起始序号或数字位数变化时的处理"""
        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                self.generate_batch_preview(self.file_list, folder_path)
        elif self.current_mode == "single" and self.path_entry.get():
            self.generate_single_preview(self.path_entry.get())

    def on_date_format_change(self, event=None):
        """日期格式变化时的处理"""
        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                self.generate_batch_preview(self.file_list, folder_path)
        elif self.current_mode == "single" and self.path_entry.get():
            self.generate_single_preview(self.path_entry.get())

    def on_ignore_existing_change(self):
        """忽略已符合格式的文件选项变化时的处理"""
        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                self.generate_batch_preview(self.file_list, folder_path)

    def is_hidden_file(self, file_path):
        """检查文件是否为隐藏文件"""
        file_name = os.path.basename(file_path)

        # 检查以点开头的文件 (Unix风格)
        if file_name.startswith('.'):
            return True

        # Windows系统检查隐藏属性
        if platform.system() == 'Windows':
            try:
                import win32api, win32con
                attrs = win32api.GetFileAttributes(file_path)
                return attrs & win32con.FILE_ATTRIBUTE_HIDDEN
            except ImportError:
                # 如果没有安装pywin32，使用os.stat
                try:
                    import stat
                    return bool(os.stat(file_path).st_file_attributes & stat.FILE_ATTRIBUTE_HIDDEN)
                except (AttributeError, ImportError):
                    pass

        return False

    def browse_folder(self):
        print("[INFO] 打开文件/文件夹选择对话框")
        if self.current_mode == "batch":
            folder_selected = filedialog.askdirectory()
            if folder_selected:
                self.path_entry.delete(0, tk.END)
                self.path_entry.insert(0, folder_selected)
                self.status_var.set(f"已选择文件夹: {os.path.basename(folder_selected)}")
                print(f"[INFO] 选择文件夹: {folder_selected}")

                # 批量模式下自动显示文件列表
                self.show_batch_files(folder_selected)
        else:
            file_selected = filedialog.askopenfilename(
                filetypes=[("所有文件", "*.*")]
            )
            if file_selected:
                self.path_entry.delete(0, tk.END)
                self.path_entry.insert(0, file_selected)
                self.status_var.set(f"已选择文件: {os.path.basename(file_selected)}")
                print(f"[INFO] 选择文件: {file_selected}")

    def browse_file(self):
        print("[INFO] 打开文件选择对话框")
        if self.current_mode == "title":
            file_selected = filedialog.askopenfilename(
                filetypes=[("Word/PDF文件", "*.docx *.pdf"),
                           ("Word文件", "*.docx"),
                           ("PDF文件", "*.pdf"),
                           ("所有文件", "*.*")]
            )
        else:
            file_selected = filedialog.askopenfilename(
                filetypes=[("所有文件", "*.*")]
            )

        if file_selected:
            self.path_entry.delete(0, tk.END)
            self.path_entry.insert(0, file_selected)
            self.status_var.set(f"已选择文件: {os.path.basename(file_selected)}")
            print(f"[INFO] 选择文件: {file_selected}")
            if self.current_mode == "single":
                self.generate_single_preview(file_selected)

    def show_batch_files(self, folder_path):
        """显示批量模式下的文件列表，过滤隐藏文件"""
        self.original_files_listbox.delete(0, tk.END)
        try:
            all_files = os.listdir(folder_path)
            # 过滤隐藏文件
            self.file_list = [
                f for f in all_files
                if os.path.isfile(os.path.join(folder_path, f))
                   and not self.is_hidden_file(os.path.join(folder_path, f))
            ]

            if not self.file_list:
                self.original_files_listbox.insert(tk.END, "所选文件夹为空")
                return

            for file in self.file_list:
                self.original_files_listbox.insert(tk.END, file)

            # 自动生成预览
            self.generate_batch_preview(self.file_list, folder_path)

        except Exception as e:
            self.original_files_listbox.insert(tk.END, f"无法读取文件夹内容: {str(e)}")

    def parse_date_format(self, format_str):
        """将用户友好的日期格式转换为Python的strftime格式"""
        format_map = {
            "YYYYMMDD": "%Y%m%d",
            "YYYY-MM-DD": "%Y-%m-%d",
            "YYMMDD": "%y%m%d",
            "YY-MM-DD": "%y-%m-%d"
        }
        return format_map.get(format_str, "%Y%m%d")

    def is_sequence_filename_format(self, filename):
        """检查文件名是否符合'序号_原文件名_日期'格式"""
        base_name, _ = os.path.splitext(filename)

        patterns = [
            r'^\d{1,9}_.+_\d{4}\d{2}\d{2}$',
            r'^\d{1,9}_.+_\d{4}-\d{2}-\d{2}$',
            r'^\d{1,9}_.+_\d{2}\d{2}\d{2}$',
            r'^\d{1,9}_.+_\d{2}-\d{2}-\d{2}$',
        ]

        for pattern in patterns:
            if re.match(pattern, base_name):
                return True

        return False

    def get_next_sequence_number(self, files):
        """获取文件夹中已有'序号_原文件名_日期'格式文件的最大序号+1"""
        max_seq = 0
        for file in files:
            if self.is_sequence_filename_format(file):
                base_name, _ = os.path.splitext(file)
                parts = base_name.split('_')
                if len(parts) >= 3:
                    try:
                        seq_num = int(parts[0])
                        if seq_num > max_seq:
                            max_seq = seq_num
                    except ValueError:
                        continue
        return max_seq + 1

    def generate_batch_preview(self, files, folder_path):
        """生成批量重命名预览"""
        self.preview_listbox.delete(0, tk.END)

        prefix = self.title_prefix_var.get().strip()
        suffix = self.title_suffix_var.get().strip()
        target_ext = self.get_target_extension()
        preset = self.preset_var.get()
        date_format = self.parse_date_format(self.date_format_var.get())
        current_date = datetime.now().strftime(date_format)
        ignore_existing = self.ignore_existing_format_var.get()

        renamed_files = []

        # 重置序号计数器
        start_num = self.start_num_var.get()
        if ignore_existing and preset == "序号_原文件名_日期":
            start_num = self.get_next_sequence_number(files)
            print(f"[INFO] 已有格式文件检测完成，起始序号设置为: {start_num}")

        seq_num = start_num
        digits = self.digits_var.get()

        for file in files:
            base_name, ext = os.path.splitext(file)

            if ignore_existing and preset == "序号_原文件名_日期" and self.is_sequence_filename_format(file):
                new_name = file
                print(f"[INFO] 保留已有格式文件: {file}")
            else:
                new_base_name = base_name
                if preset == "序号模式 (001, 002...)":
                    new_base_name = f"{seq_num:0{digits}d}"
                elif preset == "日期模式 (YYYYMMDD)":
                    new_base_name = current_date
                elif preset == "时间戳模式 (YYYYMMDD_HHMMSS)":
                    new_base_name = datetime.now().strftime("%Y%m%d_%H%M%S")
                elif preset == "序号+原文件名":
                    new_base_name = f"{seq_num:0{digits}d}_{base_name}"
                elif preset == "日期+原文件名":
                    new_base_name = f"{current_date}_{base_name}"
                elif preset == "原文件名+序号":
                    new_base_name = f"{base_name}_{seq_num:0{digits}d}"
                elif preset == "序号_原文件名_日期":
                    new_base_name = f"{seq_num:0{digits}d}_{base_name}_{current_date}"

                if prefix:
                    new_base_name = prefix + new_base_name
                if suffix:
                    new_base_name += suffix

                if target_ext:
                    new_ext = f".{target_ext}"
                else:
                    new_ext = ext

                new_name = new_base_name + new_ext

                # 只有在需要使用序号的预设模式下才增加序号
                if preset in ["序号模式 (001, 002...)", "序号+原文件名", "原文件名+序号", "序号_原文件名_日期"]:
                    seq_num += 1

            renamed_files.append((file, new_name))

            if len(renamed_files) <= 100:
                self.preview_listbox.insert(tk.END, f"{file} → {new_name}")

        if len(files) > 100:
            self.preview_listbox.insert(tk.END, f"... 共{len(files)}个文件将被重命名")

        self.preview_results = renamed_files
        self.execute_btn.config(state=tk.NORMAL if renamed_files else tk.DISABLED)
        self.remove_format_btn.config(state=tk.NORMAL if renamed_files else tk.DISABLED)
        print(f"[INFO] 生成批量重命名预览，{len(renamed_files)}个文件")

    def generate_single_preview(self, file_path):
        """生成单个文件重命名预览"""
        if not os.path.isfile(file_path):
            return

        filename = os.path.basename(file_path)
        base_name, ext = os.path.splitext(filename)
        preset = self.preset_var.get()
        prefix = self.title_prefix_var.get().strip()
        suffix = self.title_suffix_var.get().strip()
        target_ext = self.get_target_extension()
        current_date = datetime.now().strftime(self.parse_date_format(self.date_format_var.get()))
        seq_num = self.start_num_var.get()
        digits = self.digits_var.get()

        if preset == "序号模式 (001, 002...)":
            new_base_name = f"{seq_num:0{digits}d}"
        elif preset == "日期模式 (YYYYMMDD)":
            new_base_name = current_date
        elif preset == "时间戳模式 (YYYYMMDD_HHMMSS)":
            new_base_name = datetime.now().strftime("%Y%m%d_%H%M%S")
        elif preset == "序号+原文件名":
            new_base_name = f"{seq_num:0{digits}d}_{base_name}"
        elif preset == "日期+原文件名":
            new_base_name = f"{current_date}_{base_name}"
        elif preset == "原文件名+序号":
            new_base_name = f"{base_name}_{seq_num:0{digits}d}"
        elif preset == "序号_原文件名_日期":
            new_base_name = f"{seq_num:0{digits}d}_{base_name}_{current_date}"
        else:
            new_base_name = base_name

        new_base_name = f"{prefix}{new_base_name}{suffix}"
        new_name = new_base_name + (f".{target_ext}" if target_ext else ext)

        self.original_files_listbox.delete(0, tk.END)
        self.original_files_listbox.insert(tk.END, filename)
        self.preview_listbox.delete(0, tk.END)
        self.preview_listbox.insert(tk.END, f"{filename} → {new_name}")
        self.preview_results = [(filename, new_name)]
        self.execute_btn.config(state=tk.NORMAL)
        self.remove_format_btn.config(state=tk.NORMAL)

    def on_extension_change(self, event=None):
        """下拉菜单选择变化时的处理"""
        selected = self.extension_var.get()
        if selected != "不修改":
            self.custom_ext_var.set("")

        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                self.generate_batch_preview(self.file_list, folder_path)
        elif self.current_mode == "single" and self.path_entry.get():
            self.generate_single_preview(self.path_entry.get())

    def on_custom_ext_change(self, event=None):
        """自定义后缀输入变化时的处理"""
        custom_ext = self.custom_ext_var.get().strip()
        if custom_ext:
            self.extension_var.set("不修改")

        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                self.generate_batch_preview(self.file_list, folder_path)
        elif self.current_mode == "single" and self.path_entry.get():
            self.generate_single_preview(self.path_entry.get())

    def get_target_extension(self):
        """获取目标文件后缀"""
        selected_ext = self.extension_var.get()
        custom_ext = self.custom_ext_var.get().strip()

        if selected_ext != "不修改":
            return selected_ext
        elif custom_ext:
            return custom_ext.replace('.', '')
        else:
            return None

    def extract_title_from_docx(self, file_path):
        """从Word文档中提取标题"""
        try:
            print(f"[INFO] 开始从Word文档提取标题: {file_path}")
            doc = docx.Document(file_path)

            if doc.core_properties.title:
                title = doc.core_properties.title.strip()
                print(f"[INFO] 从文档属性提取标题: {title}")
                return title

            for para in doc.paragraphs[:10]:
                text = para.text.strip()
                if ("中华人民共和国" in text or "法" in text) and len(text) < 60:
                    if not re.match(r'^(第[一二三四五六七八九十0-9]+章|目?录|附件|附录|修订说明|前言|序)', text):
                        print(f"[INFO] 从正文提取标题: {text}")
                        return text[:40]

            title_candidates = []

            for para in doc.paragraphs:
                if para.style.name.startswith('Heading') or para.style.name.lower().startswith('标题'):
                    candidate = para.text.strip()
                    if len(candidate) > 5 and not re.match(r'^(第[0-9一二三四五六七八九十]+)', candidate):
                        title_candidates.append(candidate)
                        print(f"[INFO] 从标题样式提取: {candidate}")
                        break

            if not title_candidates:
                for para in doc.paragraphs[:3]:
                    text = para.text.strip()
                    if len(text) > 15 and not re.match(r'^(第[0-9一二三四五六七八九十]+|目?录|附件|附录)', text):
                        title_candidates.append(text[:40])
                        print(f"[INFO] 从正文前3段提取: {text[:40]}")
                        break

            if not title_candidates and doc.paragraphs:
                first_line = doc.paragraphs[0].text.strip()
                if len(first_line) > 5:
                    title_candidates.append(first_line[:40])
                    print(f"[INFO] 从第一行提取: {first_line[:40]}")

            return title_candidates[0] if title_candidates else "无标题"

        except Exception as e:
            error_msg = f"提取失败: {str(e)}"
            print(f"[ERROR] Word标题提取失败: {error_msg}")
            return error_msg

    def extract_title_from_pdf(self, file_path):
        """从PDF文件中提取标题"""
        try:
            print(f"[INFO] 开始从PDF提取标题: {file_path}")
            pdf = PdfReader(file_path)
            info = pdf.metadata
            if info and '/Title' in info:
                title = info['/Title']
                if isinstance(title, bytes):
                    title = title.decode('utf-8', errors='replace')
                print(f"[INFO] 从PDF元数据提取标题: {title}")
                return title.strip()

            if len(pdf.pages) > 0:
                first_page = pdf.pages[0].extract_text()
                if not first_page:
                    print("[INFO] PDF首页无文本内容")
                    return "无标题"

                lines = []
                for line in first_page.split('\n'):
                    stripped = line.strip()
                    if stripped and len(stripped) > 5:
                        lines.append(stripped)

                if lines:
                    longest_line = max(lines, key=len)
                    print(f"[INFO] 从PDF正文提取标题: {longest_line[:40]}")
                    return longest_line[:40]
                else:
                    print("[INFO] PDF首页无有效行")
                    return "无标题"

            print("[INFO] PDF无有效内容，返回无标题")
            return "无标题"
        except Exception as e:
            error_msg = f"提取失败: {str(e)}"
            print(f"[ERROR] PDF标题提取失败: {error_msg}")
            return error_msg

    def extract_title(self):
        file_path = self.path_entry.get().strip()
        if not file_path:
            messagebox.showerror("错误", "请选择文件")
            print("[ERROR] 未选择文件")
            return

        if not os.path.exists(file_path):
            messagebox.showerror("错误", "文件不存在")
            print(f"[ERROR] 文件不存在: {file_path}")
            return

        file_ext = os.path.splitext(file_path)[1].lower()
        filename = os.path.basename(file_path)

        self.original_files_listbox.delete(0, tk.END)
        self.original_files_listbox.insert(tk.END, filename)

        self.preview_listbox.delete(0, tk.END)

        try:
            self.status_var.set(f"正在提取标题: {filename}")
            print(f"[INFO] 开始提取标题流程: {filename}")

            if file_ext == '.docx':
                title = self.extract_title_from_docx(file_path)
            elif file_ext == '.pdf':
                title = self.extract_title_from_pdf(file_path)
            else:
                title = f"不支持的文件格式: {file_ext}"
                print(f"[ERROR] 不支持的文件格式: {file_ext}")

            if title.startswith("提取失败"):
                self.preview_listbox.insert(tk.END, f"错误: {title}")
                self.status_var.set("提取失败")
                self.execute_btn.config(state=tk.DISABLED)
                self.remove_format_btn.config(state=tk.DISABLED)
                print(f"[ERROR] 标题提取失败: {title}")
                return

            valid_title = re.sub(r'[\\/:*?"<>|]', '_', title)
            if not valid_title:
                valid_title = "无标题_" + datetime.now().strftime("%Y%m%d%H%M%S")
                print("[INFO] 生成默认标题: 无标题_时间戳")

            prefix = self.title_prefix_var.get().strip()
            suffix = self.title_suffix_var.get().strip()

            target_ext = self.get_target_extension()
            if target_ext:
                new_ext = f".{target_ext}"
            else:
                new_ext = os.path.splitext(filename)[1]

            preset = self.preset_var.get()
            date_format = self.parse_date_format(self.date_format_var.get())
            current_date = datetime.now().strftime(date_format)

            if preset == "序号模式 (001, 002...)":
                new_base_name = f"{self.start_num_var.get():0{self.digits_var.get()}d}"
            elif preset == "日期模式 (YYYYMMDD)":
                new_base_name = current_date
            elif preset == "时间戳模式 (YYYYMMDD_HHMMSS)":
                new_base_name = datetime.now().strftime("%Y%m%d_%H%M%S")
            elif preset == "序号+原文件名":
                new_base_name = f"{self.start_num_var.get():0{self.digits_var.get()}d}_{valid_title}"
            elif preset == "日期+原文件名":
                new_base_name = f"{current_date}_{valid_title}"
            elif preset == "原文件名+序号":
                new_base_name = f"{valid_title}_{self.start_num_var.get():0{self.digits_var.get()}d}"
            elif preset == "序号_原文件名_日期":
                new_base_name = f"{self.start_num_var.get():0{self.digits_var.get()}d}_{valid_title}_{current_date}"
            else:
                new_base_name = valid_title

            if prefix:
                new_base_name = prefix + new_base_name
            if suffix:
                new_base_name += suffix

            new_filename = new_base_name + new_ext

            self.preview_listbox.insert(tk.END, f"原始文件名: {filename}\n")
            self.preview_listbox.insert(tk.END, f"提取的标题: {title}\n")
            self.preview_listbox.insert(tk.END, f"新文件名: {new_filename}")

            self.status_var.set(f"标题提取完成: {title}")
            self.preview_results = [(filename, new_filename)]
            self.file_path = file_path

            self.execute_btn.config(state=tk.NORMAL)
            self.remove_format_btn.config(state=tk.NORMAL)
            print(f"[INFO] 标题提取成功，新文件名: {new_filename}")

            if self.auto_rename_title_var.get():
                print("[INFO] 自动重命名已启用，执行重命名")
                self.execute_rename()

        except Exception as e:
            self.preview_listbox.insert(tk.END, f"错误: {str(e)}")
            self.status_var.set(f"提取出错: {str(e)}")
            traceback.print_exc()
            self.execute_btn.config(state=tk.DISABLED)
            self.remove_format_btn.config(state=tk.DISABLED)
            print(f"[ERROR] 提取标题过程中出错: {str(e)}")

    def is_file_open(self, file_path):
        """检查文件是否被其他进程打开"""
        if not self.open_files_check_var.get():
            return False

        try:
            # 尝试以写入模式打开文件，如果失败则说明文件被占用
            with open(file_path, 'a'):
                return False
        except IOError:
            return True

    def validate_rename_plan(self, folder_path, rename_plan):
        """确认重命名计划没有重复目标或覆盖已有文件"""
        targets = {}
        errors = []

        for old_name, new_name in rename_plan:
            if old_name == new_name:
                continue

            target_key = os.path.normcase(new_name)
            if target_key in targets:
                errors.append(f"{new_name}（同时由 {targets[target_key]} 和 {old_name} 生成）")
            else:
                targets[target_key] = old_name

            if os.path.exists(os.path.join(folder_path, new_name)):
                errors.append(f"{new_name}（目标已存在）")

        if errors:
            unique_errors = list(dict.fromkeys(errors))
            messagebox.showerror("无法执行重命名", "检测到命名冲突，未执行任何操作:\n" + "\n".join(unique_errors))
            self.status_var.set("重命名已取消：命名冲突")
            return False

        return True

    def execute_rename(self):
        if not self.preview_results:
            messagebox.showinfo("提示", "没有可执行的重命名操作")
            print("[INFO] 没有可执行的重命名操作")
            return

        # 单文件和标题模式保留原有的占用文件确认；批量模式逐项跳过占用文件。
        if self.current_mode != "batch" and self.is_file_open(self.path_entry.get()):
            msg = "该文件正在被其他程序使用，重命名可能会失败。\n\n是否继续?"
            if not messagebox.askyesno("文件正在使用", msg):
                self.status_var.set("操作已取消")
                return

        try:
            if self.current_mode == "batch":
                folder_path = self.path_entry.get()
                if not os.path.isdir(folder_path):
                    messagebox.showerror("错误", "请选择有效的文件夹路径")
                    return
                if not self.validate_rename_plan(folder_path, self.preview_results):
                    return

                renamed_count = 0
                skipped_files = []
                failed_files = []

                for old_name, new_name in self.preview_results:
                    old_path = os.path.join(folder_path, old_name)
                    new_path = os.path.join(folder_path, new_name)

                    if old_name == new_name:
                        continue

                    if self.is_file_open(old_path):
                        skipped_files.append(old_name)
                        print(f"[INFO] 跳过正在使用的文件: {old_name}")
                        continue

                    try:
                        os.rename(old_path, new_path)
                        renamed_count += 1
                        print(f"[INFO] 重命名成功: {old_name} → {new_name}")
                    except Exception as e:
                        failed_files.append((old_name, str(e)))
                        print(f"[ERROR] 重命名失败: {old_name}, 错误信息: {str(e)}")

                result_lines = [
                    f"成功：{renamed_count} 个",
                    f"跳过（正在使用）：{len(skipped_files)} 个",
                    f"失败：{len(failed_files)} 个",
                ]
                if skipped_files:
                    result_lines.append("\n跳过的文件:\n" + "\n".join(skipped_files))
                if failed_files:
                    error_msg = "\n".join([f"{old_name}: {error}" for old_name, error in failed_files])
                    result_lines.append("\n失败的文件:\n" + error_msg)

                result_message = "\n".join(result_lines)
                if failed_files:
                    messagebox.showerror("批量重命名完成", result_message)
                else:
                    messagebox.showinfo("批量重命名完成", result_message)

                self.status_var.set(
                    f"重命名完成：成功{renamed_count}个，跳过{len(skipped_files)}个，失败{len(failed_files)}个"
                )
                self.show_batch_files(folder_path)
                self.preview_results = []
                self.execute_btn.config(state=tk.DISABLED)
                self.remove_format_btn.config(state=tk.DISABLED)

            else:
                old_path = self.path_entry.get()
                old_name = os.path.basename(old_path)
                new_name = self.preview_results[0][1]
                new_path = os.path.join(os.path.dirname(old_path), new_name)

                if not self.validate_rename_plan(os.path.dirname(old_path), [(old_name, new_name)]):
                    return

                try:
                    os.rename(old_path, new_path)
                    messagebox.showinfo("提示", f"文件重命名成功: {old_name} → {new_name}")
                    self.status_var.set("重命名完成")
                    self.path_entry.delete(0, tk.END)
                    self.original_files_listbox.delete(0, tk.END)
                    self.preview_listbox.delete(0, tk.END)
                    self.preview_results = []
                    self.execute_btn.config(state=tk.DISABLED)
                    self.remove_format_btn.config(state=tk.DISABLED)
                    print(f"[INFO] 重命名成功: {old_name} → {new_name}")
                except Exception as e:
                    messagebox.showerror("错误", f"文件重命名失败: {str(e)}")
                    self.status_var.set(f"重命名失败: {str(e)}")
                    print(f"[ERROR] 重命名失败: {old_name}, 错误信息: {str(e)}")

        except Exception as e:
            messagebox.showerror("错误", f"执行重命名时出错: {str(e)}")
            self.status_var.set(f"执行重命名时出错: {str(e)}")
            traceback.print_exc()
            print(f"[ERROR] 执行重命名时出错: {str(e)}")

    def remove_format(self):
        if self.current_mode == "batch":
            folder_path = self.path_entry.get()
            if not os.path.isdir(folder_path):
                messagebox.showerror("错误", "请选择有效的文件夹路径")
                return

            files = os.listdir(folder_path)
            renamed_count = 0
            failed_files = []

            for file in files:
                base_name, ext = os.path.splitext(file)

                if self.is_sequence_filename_format(file):
                    parts = base_name.split('_')
                    if len(parts) >= 3:
                        new_base_name = '_'.join(parts[1:-1])
                        new_name = new_base_name + ext
                        old_path = os.path.join(folder_path, file)
                        new_path = os.path.join(folder_path, new_name)

                        try:
                            os.rename(old_path, new_path)
                            renamed_count += 1
                            print(f"[INFO] 取消格式成功: {file} → {new_name}")
                        except Exception as e:
                            failed_files.append((file, str(e)))
                            print(f"[ERROR] 取消格式失败: {file}, 错误信息: {str(e)}")

            if renamed_count > 0:
                messagebox.showinfo("提示", f"{renamed_count}个文件取消格式成功")
            if failed_files:
                error_msg = "\n".join([f"{old_name}: {error}" for old_name, error in failed_files])
                messagebox.showerror("错误", f"以下文件取消格式失败:\n{error_msg}")

            self.status_var.set("取消格式完成")
            self.show_batch_files(folder_path)
            self.preview_results = []
            self.execute_btn.config(state=tk.DISABLED)
            self.remove_format_btn.config(state=tk.DISABLED)

        else:
            file_path = self.path_entry.get()
            if not os.path.exists(file_path):
                messagebox.showerror("错误", "文件不存在")
                return

            file_name = os.path.basename(file_path)
            base_name, ext = os.path.splitext(file_name)

            if self.is_sequence_filename_format(file_name):
                parts = base_name.split('_')
                if len(parts) >= 3:
                    new_base_name = '_'.join(parts[1:-1])
                    new_name = new_base_name + ext
                    new_path = os.path.join(os.path.dirname(file_path), new_name)

                    try:
                        os.rename(file_path, new_path)
                        messagebox.showinfo("提示", f"文件取消格式成功: {file_name} → {new_name}")
                        self.status_var.set("取消格式完成")
                        self.path_entry.delete(0, tk.END)
                        self.original_files_listbox.delete(0, tk.END)
                        self.preview_listbox.delete(0, tk.END)
                        self.preview_results = []
                        self.execute_btn.config(state=tk.DISABLED)
                        self.remove_format_btn.config(state=tk.DISABLED)
                        print(f"[INFO] 取消格式成功: {file_name} → {new_name}")
                    except Exception as e:
                        messagebox.showerror("错误", f"文件取消格式失败: {str(e)}")
                        self.status_var.set(f"取消格式失败: {str(e)}")
                        print(f"[ERROR] 取消格式失败: {file_name}, 错误信息: {str(e)}")
            else:
                messagebox.showinfo("提示", "文件不符合序号_原文件名_日期格式，无需取消格式")


if __name__ == "__main__":
    print("[INFO] 程序启动中...")

    root = tk.Tk()
    app = FileRenameApp(root)
    print("[INFO] 程序初始化完成，进入主事件循环")
    root.mainloop()

