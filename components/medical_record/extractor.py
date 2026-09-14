from __future__ import annotations

import hashlib
import random
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ALIASES = {
    "admission": ("入院日期", "入院时间", "开始时间"),
    "discharge": ("出院日期", "出院时间", "结束时间"),
    "settlement": ("结算日期", "结算时间"),
    "category": ("医疗类别名称", "医疗类别", "就诊类型"),
    "category_code": ("医疗类别编码",),
    "visit_id": ("就诊ID", "就诊id"),
    "visit_no": ("住院或门诊号", "住院号", "门诊号"),
    "name": ("人员姓名", "姓名", "患者姓名"),
    "personal_code": ("人员编号", "个人编号", "人员编码", "个人编码", "psn_no"),
    "id_card": ("证件号码", "证件号", "身份证号", "身份证号码", "身份证件号码"),
    "item_name": ("医保目录名称", "医保目录名称1", "目录名称"),
    "item_code": ("医保目录编码", "医保目录编码1", "目录编码"),
    "department": ("科室名称", "入院科室名称", "就诊科室名称"),
    "fund_payment": ("基金支付总额", "基金支付金额"),
}

SENSITIVE_PATTERNS = ("人员编号", "个人编号", "人员编码", "个人编码", "身份证", "证件号码", "证件号", "psnno", "certno")
MINIMUM_ORDER = ("医疗类别名称", "医疗类别编码", "住院或门诊号", "就诊ID", "人员姓名", "科室名称", "基金支付总额", "入院日期", "出院日期", "结算日期", "医保目录名称", "医保目录编码")
META_ORDER = ("来源文件", "来源工作表", "源行号", "归属年份", "日期来源", "标准就诊类型", "随机种子")
APP_VERSION = "1.9"
MODE_RANDOM = "按年份随机抽取"
MODE_DEPARTMENT_TOP10 = "按科室基金支付总金额前十"
EXTRACTION_MODES = (MODE_RANDOM, MODE_DEPARTMENT_TOP10)


@dataclass
class FileSummary:
    file: str
    source_rows: int = 0
    in_range_rows: int = 0
    records: int = 0
    selected: int = 0
    status: str = "成功"
    message: str = ""


class BatchExtractionError(RuntimeError):
    def __init__(self, output_path: Path, errors: list[tuple[str, str]]):
        self.output_path = output_path
        self.errors = errors
        super().__init__(f"所有输入文件均处理失败。异常报告已保存：{output_path}")


def normalize(value: Any) -> str:
    return re.sub(r"[\s_\-（）()]+", "", str(value or "")).lower()


def is_sensitive_header(header: str) -> bool:
    key = normalize(header)
    return any(normalize(pattern) in key for pattern in SENSITIVE_PATTERNS)


def parse_date(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y/%m/%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def allocate_quotas(years: Iterable[int], target: int) -> dict[int, int]:
    ordered = sorted(set(years))
    if target <= 0 or not ordered:
        return {}
    if target < len(ordered):
        return {year: 1 for year in ordered[:target]}
    base, remainder = divmod(target, len(ordered))
    quotas = {year: base for year in ordered}
    for year in ordered[-remainder:] if remainder else ():
        quotas[year] += 1
    return quotas


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _field_map(headers: list[str]) -> dict[str, str | None]:
    normalized = {normalize(header): header for header in headers if header}
    result: dict[str, str | None] = {}
    for field, aliases in ALIASES.items():
        result[field] = next((normalized[normalize(alias)] for alias in aliases if normalize(alias) in normalized), None)
    return result


def _missing_required_fields(mapping: dict[str, str | None], extraction_mode: str = MODE_RANDOM) -> list[str]:
    missing = []
    if not (mapping["admission"] or mapping["settlement"]):
        missing.append("入院时间或结算时间")
    common = (("visit_id", "就诊ID"), ("name", "姓名"))
    mode_fields = {
        MODE_RANDOM: (
            ("category", "医疗类别"),
            ("visit_no", "住院号或门诊号"),
            ("item_name", "医保目录名称"),
            ("item_code", "医保目录编码"),
        ),
        MODE_DEPARTMENT_TOP10: (
            ("department", "科室名称"),
            ("fund_payment", "基金支付总额"),
        ),
    }
    if extraction_mode not in mode_fields:
        raise ValueError(f"不支持的抽取方式：{extraction_mode}")
    for field, label in common + mode_fields[extraction_mode]:
        if not mapping[field]:
            missing.append(label)
    return missing


def _find_business_sheet(workbook, extraction_mode: str = MODE_RANDOM):
    best = None
    best_rank = (-1, -1)
    for sheet in workbook.worksheets:
        for row_no, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True), 1):
            headers = [str(value).strip() if value is not None else "" for value in row]
            mapping = _field_map(headers)
            score = sum(bool(value) for value in mapping.values())
            rank = (not _missing_required_fields(mapping, extraction_mode), score)
            if rank > best_rank:
                best = (sheet, row_no, headers, mapping)
                best_rank = rank
    if not best or not (best[3]["admission"] or best[3]["settlement"]):
        raise ValueError("未找到包含日期字段的业务工作表")
    return best


def _validate_required_fields(mapping: dict[str, str | None], extraction_mode: str = MODE_RANDOM) -> None:
    missing = _missing_required_fields(mapping, extraction_mode)
    if missing:
        raise ValueError(f"业务工作表缺少必要字段：{'、'.join(missing)}")


def _value(row: dict[str, Any], mapping: dict[str, str | None], field: str) -> Any:
    header = mapping.get(field)
    return row.get(header) if header else None


def _record_key(row: dict[str, Any], mapping: dict[str, str | None], row_no: int) -> tuple:
    fields = ("visit_id", "name", "personal_code", "id_card", "admission", "discharge")
    values = []
    for field in fields:
        value = _value(row, mapping, field)
        if field in ("admission", "discharge"):
            parsed = parse_date(value)
            values.append(parsed.isoformat(timespec="microseconds") if parsed else str(value or "").strip())
        else:
            values.append(str(value or "").strip())
    values = tuple(values)
    meaningful = any(values[index] for index in (0, 1, 2, 3))
    return values if meaningful else values + (f"__row_{row_no}",)


def _merge_items(group: list[tuple[int, dict[str, Any], datetime, str]], name_header: str, code_header: str) -> tuple[str, str]:
    pairs = []
    for entry in group:
        name = str(entry[1].get(name_header) or "").strip()
        code = str(entry[1].get(code_header) or "").strip()
        pair = (name, code)
        if pair != ("", "") and pair not in pairs:
            pairs.append(pair)
    return "；".join(name for name, _ in pairs), "；".join(code for _, code in pairs)


def read_candidates(path: Path, start: date, end: date, seed: int, extraction_mode: str = MODE_RANDOM) -> tuple[list[dict[str, Any]], FileSummary, list[str]]:
    before_hash = sha256_file(path)
    before_stat = path.stat()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet, header_row, raw_headers, mapping = _find_business_sheet(workbook, extraction_mode)
        _validate_required_fields(mapping, extraction_mode)
        headers = []
        for index, header in enumerate(raw_headers, 1):
            headers.append(header or f"未命名列{index}")
        summary = FileSummary(path.name)
        groups: dict[tuple, list[tuple[int, dict[str, Any], datetime, str]]] = defaultdict(list)
        for row_no, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not any(value not in (None, "") for value in values):
                continue
            summary.source_rows += 1
            row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
            admission = parse_date(_value(row, mapping, "admission"))
            settlement = parse_date(_value(row, mapping, "settlement"))
            decision, source = (admission, "入院时间") if admission else (settlement, "结算时间")
            if not decision or not (start <= decision.date() <= end):
                continue
            summary.in_range_rows += 1
            groups[_record_key(row, mapping, row_no)].append((row_no, row, decision, source))

        records = []
        safe_headers = [header for header in headers if not is_sensitive_header(header)]
        for group in groups.values():
            first_row = group[0][1]
            merged = {}
            for header in safe_headers:
                merged[header] = next((entry[1].get(header) for entry in group if entry[1].get(header) not in (None, "")), None)
            item_name_header = mapping.get("item_name")
            item_code_header = mapping.get("item_code")
            if item_name_header and item_code_header:
                merged[item_name_header], merged[item_code_header] = _merge_items(group, item_name_header, item_code_header)
            category = str(_value(first_row, mapping, "category") or "")
            decision = group[0][2]
            merged.update({
                "来源文件": path.name,
                "来源工作表": sheet.title,
                "源行号": "；".join(str(entry[0]) for entry in group),
                "归属年份": decision.year,
                "日期来源": group[0][3],
                "标准就诊类型": ("住院" if "住院" in category else "门诊") if extraction_mode == MODE_RANDOM else None,
                "随机种子": seed if extraction_mode == MODE_RANDOM else None,
            })
            records.append(merged)
        summary.records = len(records)
    finally:
        workbook.close()
    after_stat = path.stat()
    if before_hash != sha256_file(path) or before_stat.st_mtime_ns != after_stat.st_mtime_ns or before_stat.st_size != after_stat.st_size:
        raise RuntimeError("源文件发生变化，已中止")
    return records, summary, safe_headers


def select_records(records: list[dict[str, Any]], target: int, seed: int) -> list[dict[str, Any]]:
    by_year: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_year[int(record["归属年份"])].append(record)
    quotas = allocate_quotas(by_year.keys(), target)
    pools: dict[int, list[dict[str, Any]]] = {}
    for year, values in by_year.items():
        inpatient = [value for value in values if value["标准就诊类型"] == "住院"]
        outpatient = [value for value in values if value["标准就诊类型"] != "住院"]
        year_rng = random.Random(f"{seed}:{year}")
        year_rng.shuffle(inpatient)
        year_rng.shuffle(outpatient)
        pools[year] = inpatient + outpatient
    selected = []
    used: set[int] = set()
    for year in sorted(quotas):
        for record in pools[year][: quotas[year]]:
            selected.append(record)
            used.add(id(record))
    deficit = min(target, len(records)) - len(selected)
    if deficit > 0:
        for year in sorted(pools):
            for record in pools[year]:
                if id(record) not in used:
                    selected.append(record)
                    used.add(id(record))
                    deficit -= 1
                    if deficit == 0:
                        break
            if deficit == 0:
                break
    return selected


def _parse_amount(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        amount = Decimal(str(value).replace(",", "").strip())
        return amount if amount.is_finite() else None
    except InvalidOperation:
        return None


def select_department_top_records(records: list[dict[str, Any]], limit: int = 10) -> list[dict[str, Any]]:
    by_department: dict[str, list[tuple[Decimal, int, dict[str, Any]]]] = defaultdict(list)
    for source_order, record in enumerate(records):
        mapping = _field_map(list(record))
        department = str(_value(record, mapping, "department") or "").strip()
        amount = _parse_amount(_value(record, mapping, "fund_payment"))
        if department and amount is not None:
            by_department[department].append((amount, source_order, record))

    selected = []
    for department in sorted(by_department):
        ranked = sorted(by_department[department], key=lambda entry: (-entry[0], entry[1]))
        selected.extend(entry[2] for entry in ranked[:limit])
    return selected


def _canonical_row(record: dict[str, Any]) -> dict[str, Any]:
    mapping = _field_map(list(record))
    output = dict(record)
    canonical = {
        "医疗类别名称": _value(record, mapping, "category"),
        "医疗类别编码": _value(record, mapping, "category_code"),
        "住院或门诊号": _value(record, mapping, "visit_no"),
        "就诊ID": _value(record, mapping, "visit_id"),
        "人员姓名": _value(record, mapping, "name"),
        "科室名称": _value(record, mapping, "department"),
        "基金支付总额": _value(record, mapping, "fund_payment"),
        "入院日期": _value(record, mapping, "admission"),
        "出院日期": _value(record, mapping, "discharge"),
        "结算日期": _value(record, mapping, "settlement"),
        "医保目录名称": _value(record, mapping, "item_name"),
        "医保目录编码": _value(record, mapping, "item_code"),
    }
    for key, value in canonical.items():
        if value not in (None, ""):
            output[key] = value
    return {key: value for key, value in output.items() if not is_sensitive_header(key)}


def write_result(path: Path, records: list[dict[str, Any]], summaries: list[FileSummary], params: dict[str, Any], raw_headers: list[str], errors: list[tuple[str, str]]) -> None:
    workbook = Workbook()
    result = workbook.active
    result.title = "抽取结果"
    rows = [_canonical_row(record) for record in records]
    other_headers = []
    known_headers = {normalize(alias) for aliases in ALIASES.values() for alias in aliases}
    for header in raw_headers:
        if normalize(header) not in known_headers and header not in META_ORDER and not is_sensitive_header(header) and header not in other_headers:
            other_headers.append(header)
    headers = [header for header in MINIMUM_ORDER if any(row.get(header) not in (None, "") for row in rows)] + other_headers + list(META_ORDER)
    _write_sheet(result, headers, [[row.get(header) for header in headers] for row in rows])

    summary_sheet = workbook.create_sheet("运行汇总")
    summary_headers = ["源文件", "源数据行数", "检查范围内行数", "范围外或日期无效行数", "合并后病历数", "实际抽取数", "状态", "说明"]
    _write_sheet(summary_sheet, summary_headers, [[s.file, s.source_rows, s.in_range_rows, s.source_rows - s.in_range_rows, s.records, s.selected, s.status, s.message] for s in summaries])

    error_sheet = workbook.create_sheet("异常明细")
    _write_sheet(error_sheet, ["源文件", "异常说明"], errors)

    param_sheet = workbook.create_sheet("参数")
    param_rows = [["参数", "值"], *[[key, value] for key, value in params.items()]]
    for row in param_rows:
        param_sheet.append(row)
    _style_sheet(param_sheet)
    workbook.save(path)


def _write_sheet(sheet, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    _style_sheet(sheet)


def _style_sheet(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    if sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 100) + 1)]
        sheet.column_dimensions[get_column_letter(column)].width = min(max(max(map(len, values), default=8) + 2, 10), 32)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (datetime, date)):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"


def unique_output_path(directory: Path, seed: int) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = directory / f"病历抽取结果_{stamp}_种子{seed}.xlsx"
    if not base.exists():
        return base
    for index in range(1, 1000):
        candidate = directory / f"病历抽取结果_{stamp}_种子{seed}_{index}.xlsx"
        if not candidate.exists():
            return candidate
    raise RuntimeError("无法生成不重名的输出文件")


def extract_files(files: Iterable[str | Path], start: date, end: date, target: int, seed: int, output_dir: str | Path, return_errors: bool = False, extraction_mode: str = MODE_RANDOM):
    if start > end:
        raise ValueError("检查开始日期不能晚于结束日期")
    if extraction_mode not in EXTRACTION_MODES:
        raise ValueError(f"不支持的抽取方式：{extraction_mode}")
    if extraction_mode == MODE_RANDOM and target < 1:
        raise ValueError("每文件抽取条数必须大于 0")
    sources = [Path(file).resolve() for file in files]
    if not sources:
        raise ValueError("请至少选择一个输入文件")
    records, summaries, raw_headers, errors = [], [], [], []
    for source in sources:
        try:
            file_seed = int(hashlib.sha256(f"{seed}|{source.name.lower()}".encode("utf-8")).hexdigest()[:16], 16)
            candidates, summary, headers = read_candidates(source, start, end, file_seed, extraction_mode)
            if extraction_mode == MODE_DEPARTMENT_TOP10:
                selected = select_department_top_records(candidates)
            else:
                selected = select_records(candidates, target, file_seed)
            summary.selected = len(selected)
            if extraction_mode == MODE_RANDOM and len(selected) < target:
                summary.message = f"候选不足：目标 {target}，实际 {len(selected)}"
            records.extend(selected)
            summaries.append(summary)
            raw_headers.extend(header for header in headers if header not in raw_headers)
        except Exception as exc:
            summaries.append(FileSummary(source.name, status="失败", message=str(exc)))
            errors.append((source.name, str(exc)))
    output = unique_output_path(Path(output_dir).resolve(), seed)
    if output in sources:
        raise ValueError("输出文件不能与源文件相同")
    params = {
        "工具版本": APP_VERSION,
        "抽取方式": extraction_mode,
        "检查开始日期": start.isoformat(),
        "检查结束日期": end.isoformat(),
        "每文件目标条数": target if extraction_mode == MODE_RANDOM else "不适用（每科室固定前10）",
        "随机种子": seed if extraction_mode == MODE_RANDOM else "不适用",
        "隐私规则": "姓名完整保留；个人/人员编号、身份证/证件号整列不输出",
    }
    write_result(output, records, summaries, params, raw_headers, errors)
    if len(errors) == len(sources):
        raise BatchExtractionError(output, errors)
    return (output, errors) if return_errors else output
