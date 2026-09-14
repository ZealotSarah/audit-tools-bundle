from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from audit_tools_bundle.tabs.medical_record import load_settings, save_settings
from audit_tools_bundle.workers import medical_record as medical_worker
from components.medical_record.extractor import (
    MODE_DEPARTMENT_TOP10,
    MODE_RANDOM,
    extract_files,
    read_candidates,
    select_department_top_records,
    sha256_file,
)


class Messages:
    def __init__(self):
        self.items = []

    def put(self, value):
        self.items.append(value)


class NotCancelled:
    def is_set(self):
        return False


def test_department_top_ten_ranks_each_department_and_keeps_source_unchanged(tmp_path):
    source = tmp_path / "sql-export.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["就诊ID", "人员编号", "证件号码", "人员姓名", "开始时间", "科室名称", "基金支付总额"])
    for department in ("内科", "外科"):
        for index in range(12):
            sheet.append([
                f"{department}-{index}", f"P{index}", f"ID{index}", f"患者{index}",
                datetime(2025, 1, index + 1), department, f"{index * 100:,}.50",
            ])
    sheet.append(["NO-DEPT", "P99", "ID99", "无科室", datetime(2025, 1, 1), None, 999999])
    sheet.append(["BAD-AMOUNT", "P98", "ID98", "坏金额", datetime(2025, 1, 1), "内科", "无法解析"])
    workbook.save(source)
    workbook.close()
    source_hash = sha256_file(source)

    output = extract_files(
        [source], date(2025, 1, 1), date(2025, 12, 31), 0, 0, tmp_path / "out",
        extraction_mode=MODE_DEPARTMENT_TOP10,
    )

    assert sha256_file(source) == source_hash
    result = load_workbook(output, data_only=True)
    result_sheet = result["抽取结果"]
    headers = [cell.value for cell in result_sheet[1]]
    rows = [dict(zip(headers, row)) for row in result_sheet.iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 20
    assert {row["科室名称"] for row in rows} == {"内科", "外科"}
    for department in ("内科", "外科"):
        department_rows = [row for row in rows if row["科室名称"] == department]
        assert [float(str(row["基金支付总额"]).replace(",", "")) for row in department_rows] == [
            index * 100 + 0.5 for index in range(11, 1, -1)
        ]
    assert "人员编号" not in headers
    assert "证件号码" not in headers
    params = dict(result["参数"].iter_rows(min_row=2, values_only=True))
    assert params["抽取方式"] == MODE_DEPARTMENT_TOP10
    assert params["每文件目标条数"] == "不适用（每科室固定前10）"
    result.close()


def test_department_mode_requires_ranking_columns(tmp_path):
    source = tmp_path / "missing-ranking-fields.xlsx"
    workbook = Workbook()
    workbook.active.append(["就诊ID", "人员姓名", "开始时间"])
    workbook.active.append(["V1", "张三", datetime(2025, 1, 1)])
    workbook.save(source)
    workbook.close()

    with pytest.raises(ValueError, match="科室名称") as caught:
        read_candidates(source, date(2025, 1, 1), date(2025, 12, 31), 1, MODE_DEPARTMENT_TOP10)
    assert "基金支付总额" in str(caught.value)


def test_department_ranking_ties_are_stable_and_invalid_amounts_are_skipped():
    records = [
        {"科室名称": "内科", "基金支付总额": "100", "就诊ID": "first"},
        {"科室名称": "内科", "基金支付总额": "NaN", "就诊ID": "invalid"},
        {"科室名称": "内科", "基金支付总额": 100, "就诊ID": "second"},
    ]
    assert [record["就诊ID"] for record in select_department_top_records(records)] == ["first", "second"]


def test_medical_settings_remember_mode_and_default_legacy_settings(tmp_path):
    path = tmp_path / "settings.json"
    save_settings(
        date(2025, 1, 1), date(2025, 12, 31), "C:/input", "D:/output",
        MODE_DEPARTMENT_TOP10, path,
    )
    assert load_settings(path)["extraction_mode"] == MODE_DEPARTMENT_TOP10

    path.write_text('{"start_date":"2025-01-01","end_date":"2025-12-31"}', encoding="utf-8")
    assert load_settings(path)["extraction_mode"] == MODE_RANDOM

    path.write_text('{"start_date":"2025-01-01","end_date":"2025-12-31","extraction_mode":"bad"}', encoding="utf-8")
    assert load_settings(path) == {}


def test_medical_worker_forwards_mode_and_defaults_old_payload(monkeypatch, tmp_path):
    modes = []

    def fake_extract_files(*_args, extraction_mode, **_kwargs):
        modes.append(extraction_mode)
        return tmp_path / "result.xlsx", []

    monkeypatch.setattr(medical_worker, "extract_files", fake_extract_files)
    payload = {
        "files": ["input.xlsx"], "start": "2025-01-01", "end": "2025-12-31",
        "count": 0, "seed": 0, "output_dir": str(tmp_path),
    }
    medical_worker.run_medical_extraction(Messages(), NotCancelled(), {
        **payload, "extraction_mode": MODE_DEPARTMENT_TOP10,
    })
    medical_worker.run_medical_extraction(Messages(), NotCancelled(), payload)

    assert modes == [MODE_DEPARTMENT_TOP10, MODE_RANDOM]
