from decimal import Decimal

from openpyxl import Workbook, load_workbook

from components.fund_calculator.fund_calculator import APP_VERSION, RunOptions, run_file


PAIR_HEADERS = [
    "定点编码", "定点名称", "就诊ID", "医疗类别编码", "医疗类别", "险种类型",
    "医疗费总额", "基金支付总额",
    "医保目录编码1", "医保目录名称1", "单价1", "处方日期1", "数量1",
    "医保目录编码2", "医保目录名称2", "单价2", "处方日期2", "数量2",
]


def test_fund_component_version_is_1_0_0():
    assert APP_VERSION == "1.0.0"


def test_simultaneous_pair_is_available_in_integrated_core(tmp_path):
    path = tmp_path / "pair.xlsx"
    workbook = Workbook()
    source = workbook.active
    source.append(PAIR_HEADERS)
    source.append([
        "H1", "医院", "V1", 21, "普通住院", "城乡", 1000, 500,
        "A", "项目一", 10, "2026-09-19 10:15:03", 3,
        "B", "项目二", 20, "2026-09-19 10:15:48", 2,
    ])
    workbook.save(path)
    workbook.close()

    result = run_file(path, RunOptions(
        "两项同时收取", "自动识别", "廊坊市", "三级",
        None, None, True, None, "同日同时同分",
    ))

    assert (result.processed, result.successful, result.errors) == (1, 1, 0)
    assert result.total_fund == Decimal("34.10")
    saved = load_workbook(path, data_only=True)
    try:
        output = saved["基金测算"]
        assert [output.cell(7, column).value for column in range(1, 10)] == [
            "医疗机构编码", "医疗机构名称", "险种类别", "医疗类别", "医疗总额",
            "数量总和", "人次", "基金金额", "报销比例",
        ]
        assert output["F8"].value == 4
        assert output["G8"].value == 1
        assert output["H8"].value == 34.1
        assert output["K7"].value == "同日同时同分"
    finally:
        saved.close()
